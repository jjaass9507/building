import hashlib
import json
import re
import uuid
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, Iterable, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SCHEMA_VERSION = "1.0"
BUILDING_ID_KEY = "_building_id"
FLOOR_ID_KEY = "_floor_id"

BUILDING_FIELDS = [
    "棟別",
    "基地面積(M2)",
    "容積率",
    "建蔽率",
    "開挖深度(M)",
    "耐震係數(gal)",
    "汽車停車位",
    "機車停車位",
]

FLOOR_FIELDS = [
    "棟別",
    "樓層",
    "狀態",
    "預計成廠年份",
    "進駐製程",
    "樓層高度(M)",
    "無塵室淨高(M)",
    "樓地板面積(M2)",
    "無塵室面積(M2)",
    "生產週邊(M2)",
    "公設(含其他)(公式)(M2)",
    "廠務設施面積(M2)",
    "純水",
    "廢水",
    "給排水",
    "空調",
    "抽氣",
    "氣體",
    "電力",
    "弱電",
    "消防",
    "監控",
    "監控/弱電/消防",
    "樓層載重kgf/m2",
]

READABLE_HEADERS = FLOOR_FIELDS + [field for field in BUILDING_FIELDS if field != "棟別"]

NUMERIC_BUILDING_FIELDS = set(BUILDING_FIELDS[1:])
NUMERIC_FLOOR_FIELDS = {
    "樓層高度(cm)",
    "無塵室淨高(cm)",
    "樓地板面積(M2)",
    "無塵室面積(M2)",
    "生產週邊(M2)",
    "公設(含其他)(公式)(M2)",
    "樓層載重kgf/m2",
}
FACILITY_DETAIL_FIELDS = [
    "純水", "廢水", "給排水", "空調", "抽氣", "氣體", "電力", "弱電", "消防", "監控", "其他"
]

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
HEADER_FILL = PatternFill("solid", fgColor="334155")
SUBHEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
ACCENT_FILL = PatternFill("solid", fgColor="DBEAFE")
CHANGE_TYPE_LABELS = {
    "ADD": "新增",
    "ADJUST": "資料修正",
    "EXPAND": "擴建",
    "REDUCE": "面積減少",
    "DEMOLISH": "拆除",
    "IMPORT": "Excel 匯入",
}


class BuildingDataError(ValueError):
    pass


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any, field: str) -> float:
    if value in (None, ""):
        return 0.0
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise BuildingDataError(f"欄位「{field}」必須是數字。") from exc
    if number != number or number in (float("inf"), float("-inf")):
        raise BuildingDataError(f"欄位「{field}」必須是有限數字。")
    return number


_HEIGHT_VALUE_PATTERN = re.compile(
    r"^([+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*(cm|公分|m|公尺|meter|meters)?$",
    re.IGNORECASE,
)
_HEIGHT_NUMBER_PATTERN = re.compile(
    r"([+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*(cm|公分|m|公尺|meter|meters)?",
    re.IGNORECASE,
)
_HEIGHT_EMPTY_MARKERS = {
    "-", "--", "—", "－", "n/a", "na", "none", "null", "無", "未設", "未提供", "未定", "待確認", "不適用", "tbd",
}


def _height_cm(value: Any, field: str) -> float:
    """將高度統一為 cm，並相容早期以 M／cm 字串保存的資料。"""
    if value in (None, ""):
        return 0.0
    if isinstance(value, dict) and "value" in value:
        value = value["value"]

    if isinstance(value, str):
        text = value.strip().replace("，", ",")
        if not text or text.casefold() in _HEIGHT_EMPTY_MARKERS:
            return 0.0

        # 同一樓層可能記錄多個區域高度，例如「2.50；2.80」。
        # 單一數值欄位採較低的淨高，避免匯出後高估樓層可用高度。
        parts = [part.strip() for part in re.split(r"[;；/／~～]", text) if part.strip()]
        if len(parts) > 1:
            return min(_height_cm(part, field) for part in parts)

        match = _HEIGHT_VALUE_PATTERN.fullmatch(text)
        if not match:
            # 允許「廠務機房：7.70停車空間：3.60」等區域標註格式。
            # 擷取全部高度後採最低值，作為該樓層保守的可用淨高。
            labelled_values = _HEIGHT_NUMBER_PATTERN.findall(text)
            if labelled_values:
                return min(
                    _height_cm(f"{number}{unit or ''}", field)
                    for number, unit in labelled_values
                )
            raise BuildingDataError(
                f"欄位「{field}」必須包含可辨識的高度數字；目前值為 {value!r}。"
            )
        number = _number(match.group(1).replace(",", ""), field)
        unit = (match.group(2) or "").lower()
        if unit in {"m", "公尺", "meter", "meters"}:
            return number * 100
        if unit in {"cm", "公分"}:
            return number
    else:
        number = _number(value, field)

    # 舊版 Excel 的欄名為 M，但資料曾直接存進 cm 欄位；高度小於 100
    # 且未標示單位時視為公尺，避免匯出時誤顯示成 0.035 M。
    if number and abs(number) < 100:
        return number * 100
    return number


def _stable_id(prefix: str, *parts: str) -> str:
    value = "|".join(_text(part).lower() for part in parts)
    return f"{prefix}-{uuid.uuid5(uuid.NAMESPACE_URL, value)}"


def _facility_value(value: Any) -> Tuple[float, Dict[str, float]]:
    if isinstance(value, dict):
        details = {
            _text(key): _number(item, _text(key))
            for key, item in (value.get("details") or {}).items()
            if _text(key)
        }
        return _number(value.get("value"), "廠務設施面積(M2)"), details
    return _number(value, "廠務設施面積(M2)"), {}


def normalize_dataset(data: Any) -> List[Dict[str, Any]]:
    """驗證並正規化平台使用的巢狀建物資料。"""
    if not isinstance(data, list):
        raise BuildingDataError("建物資料必須是陣列。")
    if len(data) > 2000:
        raise BuildingDataError("建物筆數超過系統上限 2,000 筆。")

    normalized: List[Dict[str, Any]] = []
    building_names = set()
    building_ids = set()

    for building_index, source_building in enumerate(data, start=1):
        if not isinstance(source_building, dict):
            raise BuildingDataError(f"第 {building_index} 筆建物資料格式錯誤。")

        building_name = _text(source_building.get("棟別"))
        if not building_name:
            raise BuildingDataError(f"第 {building_index} 筆建物未填寫棟別。")
        name_key = building_name.casefold()
        if name_key in building_names:
            raise BuildingDataError(f"棟別「{building_name}」重複。")
        building_names.add(name_key)

        building_id = _text(source_building.get(BUILDING_ID_KEY)) or _stable_id("BLD", building_name)
        if building_id in building_ids:
            raise BuildingDataError(f"棟別「{building_name}」的識別碼重複。")
        building_ids.add(building_id)

        building: Dict[str, Any] = {
            BUILDING_ID_KEY: building_id,
            "棟別": building_name,
        }
        for field in NUMERIC_BUILDING_FIELDS:
            building[field] = _number(source_building.get(field), field)

        floors = source_building.get("樓層") or []
        if not isinstance(floors, list):
            raise BuildingDataError(f"棟別「{building_name}」的樓層資料必須是陣列。")
        if len(floors) > 500:
            raise BuildingDataError(f"棟別「{building_name}」的樓層筆數超過 500 筆。")

        normalized_floors: List[Dict[str, Any]] = []
        floor_names = set()
        floor_ids = set()
        for floor_index, source_floor in enumerate(floors, start=1):
            if not isinstance(source_floor, dict):
                raise BuildingDataError(f"棟別「{building_name}」第 {floor_index} 筆樓層格式錯誤。")

            floor_name = _text(source_floor.get("樓層"))
            if not floor_name:
                raise BuildingDataError(f"棟別「{building_name}」第 {floor_index} 筆未填寫樓層。")
            floor_key = floor_name.casefold()
            if floor_key in floor_names:
                raise BuildingDataError(f"棟別「{building_name}」的樓層「{floor_name}」重複。")
            floor_names.add(floor_key)

            floor_id = _text(source_floor.get(FLOOR_ID_KEY)) or _stable_id("FLR", building_id, floor_name)
            if floor_id in floor_ids:
                raise BuildingDataError(f"棟別「{building_name}」的樓層識別碼重複。")
            floor_ids.add(floor_id)

            floor: Dict[str, Any] = {
                FLOOR_ID_KEY: floor_id,
                "樓層": floor_name,
                "狀態": _text(source_floor.get("狀態")) or "已成廠",
                "預計成廠年份": _text(source_floor.get("預計成廠年份")),
                "進駐製程": _text(source_floor.get("進駐製程")),
            }
            for field in NUMERIC_FLOOR_FIELDS:
                normalizer = _height_cm if field in {"樓層高度(cm)", "無塵室淨高(cm)"} else _number
                field_label = f"棟別「{building_name}」樓層「{floor_name}」的欄位「{field}」"
                floor[field] = normalizer(source_floor.get(field), field_label)

            facility_value, facility_details = _facility_value(source_floor.get("廠務設施面積(M2)"))
            floor["廠務設施面積(M2)"] = {
                "value": facility_value,
                "details": facility_details,
            }
            building["樓層"] = normalized_floors
            normalized_floors.append(floor)

        building["樓層"] = normalized_floors
        normalized.append(building)

    return normalized


def dataset_revision(data: Any) -> str:
    serialized = json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


def dataset_counts(data: Iterable[Dict[str, Any]]) -> Dict[str, int]:
    items = list(data)
    return {
        "buildings": len(items),
        "floors": sum(len(item.get("樓層") or []) for item in items),
    }


def summarize_changes(before: List[Dict[str, Any]], after: List[Dict[str, Any]]) -> Dict[str, Any]:
    before_buildings = {item.get(BUILDING_ID_KEY): item for item in before}
    after_buildings = {item.get(BUILDING_ID_KEY): item for item in after}
    before_floors = {
        floor.get(FLOOR_ID_KEY): floor
        for building in before
        for floor in building.get("樓層") or []
    }
    after_floors = {
        floor.get(FLOOR_ID_KEY): floor
        for building in after
        for floor in building.get("樓層") or []
    }

    common_buildings = before_buildings.keys() & after_buildings.keys()
    common_floors = before_floors.keys() & after_floors.keys()
    before_area = sum(
        _number(floor.get("樓地板面積(M2)"), "樓地板面積(M2)")
        for building in before
        for floor in building.get("樓層") or []
    )
    after_area = sum(
        _number(floor.get("樓地板面積(M2)"), "樓地板面積(M2)")
        for building in after
        for floor in building.get("樓層") or []
    )
    return {
        "buildings_added": len(after_buildings.keys() - before_buildings.keys()),
        "buildings_removed": len(before_buildings.keys() - after_buildings.keys()),
        "buildings_updated": sum(before_buildings[key] != after_buildings[key] for key in common_buildings),
        "floors_added": len(after_floors.keys() - before_floors.keys()),
        "floors_removed": len(before_floors.keys() - after_floors.keys()),
        "floors_updated": sum(before_floors[key] != after_floors[key] for key in common_floors),
        "area_delta_m2": round(after_area - before_area, 4),
    }


def append_audit_record(path: str, record: Dict[str, Any], max_records: int = 1000) -> None:
    records = load_audit_records(path)
    records.append(record)
    records = records[-max_records:]
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(records, handle, ensure_ascii=False, indent=2)


def load_audit_records(path: str) -> List[Dict[str, Any]]:
    try:
        with open(path, "r", encoding="utf-8") as handle:
            data = json.load(handle)
        return data if isinstance(data, list) else []
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        return []


def _display_number(value: float) -> Any:
    return int(value) if float(value).is_integer() else value


def _original_floor_value(floor: Dict[str, Any], header: str) -> Any:
    if header == "樓層高度(M)":
        return _display_number(_number(floor.get("樓層高度(cm)"), "樓層高度(cm)") / 100)
    if header == "無塵室淨高(M)":
        return _display_number(_number(floor.get("無塵室淨高(cm)"), "無塵室淨高(cm)") / 100)
    if header == "廠務設施面積(M2)":
        return _display_number(_facility_value(floor.get(header))[0])
    if header in FACILITY_DETAIL_FIELDS:
        return _display_number(_facility_value(floor.get("廠務設施面積(M2)"))[1].get(header, 0.0))
    if header == "監控/弱電/消防":
        return _display_number(_facility_value(floor.get("廠務設施面積(M2)"))[1].get("其他", 0.0))
    return floor.get(header, "")


def _flatten_readable_rows(data: List[Dict[str, Any]]) -> List[Tuple[Dict[str, Any], Dict[str, Any]]]:
    rows = []
    for building in data:
        floors = building.get("樓層") or []
        if not floors:
            floors = [{"樓層": ""}]
        rows.extend((building, floor) for floor in floors)
    return rows


def _style_title_sheet(sheet, title: str, headers: List[str], rows: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(1, 1, title)
    title_cell.font = Font(size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1E3A5F")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 30

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    sheet.row_dimensions[2].height = 42
    sheet.freeze_panes = "C3"
    if rows:
        sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{rows + 2}"


def _fit_columns(sheet, headers: List[str], max_width: int = 28) -> None:
    for index, header in enumerate(headers, start=1):
        values = [sheet.cell(row, index).value for row in range(1, min(sheet.max_row, 80) + 1)]
        width = max([len(str(header))] + [len(str(value)) for value in values if value is not None]) + 2
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width, 10), max_width)


def _write_simple_sheet(workbook: Workbook, title: str, headers: List[str], rows: List[List[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    _fit_columns(sheet, headers)


def _parse_year(value: Any) -> Tuple[int, str]:
    text = _text(value).upper()
    if not text:
        return 0, "現況"
    match = re.search(r"Y\s*(\d{1,4})", text, re.IGNORECASE) or re.search(r"(\d{4})", text)
    if not match:
        return 0, "現況"
    year = int(match.group(1))
    return year, f"Y{year}" if len(match.group(1)) < 4 else str(year)


def _annual_rows(data: List[Dict[str, Any]]) -> List[List[Any]]:
    groups: Dict[int, Dict[str, Any]] = {}
    for building in data:
        floor_groups: Dict[int, float] = {}
        labels: Dict[int, str] = {}
        for floor in building.get("樓層") or []:
            year, label = _parse_year(floor.get("預計成廠年份"))
            floor_groups[year] = floor_groups.get(year, 0.0) + _number(floor.get("樓地板面積(M2)"), "樓地板面積(M2)")
            labels[year] = label
        for year, area in floor_groups.items():
            item = groups.setdefault(year, {"label": labels[year], "buildings": set(), "area": 0.0})
            item["buildings"].add(building.get("棟別"))
            item["area"] += area

    rows = []
    cumulative = 0.0
    for year in sorted(groups):
        item = groups[year]
        starting = cumulative
        cumulative += item["area"]
        rate = item["area"] / starting if starting > 0 else None
        rows.append([
            item["label"],
            _display_number(starting),
            "、".join(sorted(item["buildings"])),
            _display_number(item["area"]),
            0,
            _display_number(cumulative),
            rate,
        ])
    return rows


def build_readable_workbook(
    data: List[Dict[str, Any]],
    audit_records: List[Dict[str, Any]],
    exported_by: str,
) -> BytesIO:
    normalized = normalize_dataset(deepcopy(data))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "建物面積總表"
    flattened = _flatten_readable_rows(normalized)
    _style_title_sheet(sheet, "建物面積資訊", READABLE_HEADERS, len(flattened))

    numeric_headers = set(READABLE_HEADERS) - {"棟別", "樓層", "狀態", "預計成廠年份", "進駐製程"}
    current_row = 3
    for building_index, building in enumerate(normalized):
        floors = building.get("樓層") or [{"樓層": ""}]
        fill = PatternFill("solid", fgColor="F8FAFC" if building_index % 2 == 0 else "EFF6FF")
        start_row = current_row
        for floor in floors:
            for column, header in enumerate(READABLE_HEADERS, start=1):
                if header in BUILDING_FIELDS:
                    value = building.get(header, "")
                elif header == "棟別":
                    value = building.get("棟別", "")
                else:
                    value = _original_floor_value(floor, header)
                cell = sheet.cell(current_row, column, value)
                cell.fill = fill
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if header in numeric_headers:
                    cell.number_format = "#,##0.00"
            current_row += 1

        end_row = current_row - 1
        if end_row > start_row:
            for header in BUILDING_FIELDS:
                column = READABLE_HEADERS.index(header) + 1
                sheet.merge_cells(start_row=start_row, start_column=column, end_row=end_row, end_column=column)
                sheet.cell(start_row, column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    _fit_columns(sheet, READABLE_HEADERS, max_width=24)
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 10

    annual_headers = ["年份", "期初面積(M2)", "新增建物", "新增面積(M2)", "減少面積(M2)", "期末累積面積(M2)", "年增率"]
    annual_rows = _annual_rows(normalized)
    _write_simple_sheet(workbook, "年度新增明細", annual_headers, annual_rows)
    annual_sheet = workbook["年度新增明細"]
    for cell in annual_sheet[1]:
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for row in range(2, annual_sheet.max_row + 1):
        annual_sheet.cell(row, 7).number_format = "0.0%"

    audit_headers = ["異動時間", "生效日期", "異動類型", "維護人員", "異動原因", "資料來源", "面積淨異動(M2)", "新增建物", "刪除建物", "修改建物", "新增樓層", "刪除樓層", "修改樓層"]
    audit_rows = []
    for record in reversed(audit_records):
        summary = record.get("summary") or {}
        audit_rows.append([
            record.get("changed_at", ""), record.get("effective_date", ""), CHANGE_TYPE_LABELS.get(record.get("change_type"), record.get("change_type", "")),
            record.get("changed_by", ""), record.get("reason", ""), record.get("source_reference", ""), summary.get("area_delta_m2", 0),
            summary.get("buildings_added", 0), summary.get("buildings_removed", 0), summary.get("buildings_updated", 0),
            summary.get("floors_added", 0), summary.get("floors_removed", 0), summary.get("floors_updated", 0),
        ])
    _write_simple_sheet(workbook, "異動紀錄", audit_headers, audit_rows)

    counts = dataset_counts(normalized)
    _write_simple_sheet(workbook, "資料說明", ["項目", "內容"], [
        ["格式版本", SCHEMA_VERSION],
        ["匯出時間", datetime.now().isoformat(timespec="seconds")],
        ["匯出人員", exported_by],
        ["面積單位", "平方公尺（M2）"],
        ["建物數", counts["buildings"]],
        ["樓層數", counts["floors"]],
        ["使用說明", "第一張工作表沿用原始 Input Excel 欄位與標題列位置，可供人員閱讀，也可重新上傳平台。"],
    ])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_standard_workbook(
    data: List[Dict[str, Any]],
    exported_by: str,
    audit_records: List[Dict[str, Any]] = None,
) -> BytesIO:
    normalized = normalize_dataset(deepcopy(data))
    workbook = Workbook()
    workbook.remove(workbook.active)
    counts = dataset_counts(normalized)

    _write_simple_sheet(workbook, "metadata", ["key", "value"], [
        ["schema_version", SCHEMA_VERSION],
        ["exported_at", datetime.now().isoformat(timespec="seconds")],
        ["exported_by", exported_by],
        ["area_unit", "m2"],
        ["building_count", counts["buildings"]],
        ["floor_count", counts["floors"]],
    ])

    building_headers = [
        "building_id", "building_code", "site_area_m2", "floor_area_ratio", "building_coverage_ratio",
        "excavation_depth_m", "seismic_coefficient_gal", "car_parking_spaces", "motorcycle_parking_spaces",
    ]
    building_rows = [[
        building.get(BUILDING_ID_KEY), building.get("棟別"), building.get("基地面積(M2)"), building.get("容積率"),
        building.get("建蔽率"), building.get("開挖深度(M)"), building.get("耐震係數(gal)"),
        building.get("汽車停車位"), building.get("機車停車位"),
    ] for building in normalized]
    _write_simple_sheet(workbook, "building_master", building_headers, building_rows)

    floor_headers = [
        "floor_id", "building_id", "building_code", "floor_name", "status", "expected_completion_year", "process_name",
        "floor_height_cm", "cleanroom_clear_height_cm", "floor_area_m2", "cleanroom_area_m2", "production_support_area_m2",
        "public_area_m2", "facility_area_m2", "facility_pure_water_m2", "facility_wastewater_m2", "facility_plumbing_m2",
        "facility_hvac_m2", "facility_exhaust_m2", "facility_gas_m2", "facility_power_m2", "facility_low_voltage_m2",
        "facility_fire_m2", "facility_monitoring_m2", "facility_other_m2", "floor_load_kgf_m2",
    ]
    floor_rows = []
    for building in normalized:
        for floor in building.get("樓層") or []:
            facility_value, details = _facility_value(floor.get("廠務設施面積(M2)"))
            floor_rows.append([
                floor.get(FLOOR_ID_KEY), building.get(BUILDING_ID_KEY), building.get("棟別"), floor.get("樓層"),
                floor.get("狀態"), floor.get("預計成廠年份"), floor.get("進駐製程"), floor.get("樓層高度(cm)"),
                floor.get("無塵室淨高(cm)"), floor.get("樓地板面積(M2)"), floor.get("無塵室面積(M2)"),
                floor.get("生產週邊(M2)"), floor.get("公設(含其他)(公式)(M2)"), facility_value,
                details.get("純水", 0), details.get("廢水", 0), details.get("給排水", 0), details.get("空調", 0),
                details.get("抽氣", 0), details.get("氣體", 0), details.get("電力", 0), details.get("弱電", 0),
                details.get("消防", 0), details.get("監控", 0), details.get("其他", 0), floor.get("樓層載重kgf/m2"),
            ])
    _write_simple_sheet(workbook, "floor_area_detail", floor_headers, floor_rows)

    change_headers = [
        "changed_at", "effective_date", "change_type", "changed_by", "reason", "source_reference", "area_delta_m2",
        "buildings_added", "buildings_removed", "buildings_updated", "floors_added", "floors_removed", "floors_updated",
    ]
    change_rows = []
    for record in audit_records or []:
        summary = record.get("summary") or {}
        change_rows.append([
            record.get("changed_at", ""), record.get("effective_date", ""), record.get("change_type", ""),
            record.get("changed_by", ""), record.get("reason", ""), record.get("source_reference", ""), summary.get("area_delta_m2", 0),
            summary.get("buildings_added", 0), summary.get("buildings_removed", 0), summary.get("buildings_updated", 0),
            summary.get("floors_added", 0), summary.get("floors_removed", 0), summary.get("floors_updated", 0),
        ])
    _write_simple_sheet(workbook, "change_log", change_headers, change_rows)

    dictionary_rows = [
        ["building_master", "building_id", "建物穩定識別碼", "string", "required"],
        ["building_master", "building_code", "棟別名稱／代碼", "string", "required"],
        ["building_master", "site_area_m2", "基地面積", "number", "m2"],
        ["floor_area_detail", "floor_id", "樓層穩定識別碼", "string", "required"],
        ["floor_area_detail", "building_id", "所屬建物識別碼", "string", "required"],
        ["floor_area_detail", "floor_area_m2", "樓地板面積", "number", "m2"],
        ["floor_area_detail", "expected_completion_year", "預計成廠年份", "string", "現況、Y1 或西元年"],
        ["floor_area_detail", "status", "資料狀態", "string", "已成廠／未成廠"],
        ["change_log", "effective_date", "資料異動生效日期", "date", "YYYY-MM-DD"],
        ["change_log", "change_type", "資料異動類型", "string", "ADD／ADJUST／EXPAND／REDUCE／DEMOLISH"],
    ]
    _write_simple_sheet(workbook, "data_dictionary", ["sheet", "field", "description", "type", "rule_or_unit"], dictionary_rows)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
