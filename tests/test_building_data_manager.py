import json
import os
import tempfile
import unittest
from copy import deepcopy

from openpyxl import load_workbook

from building_data_manager import (
    build_readable_workbook,
    build_standard_workbook,
    dataset_revision,
    normalize_dataset,
    summarize_changes,
)
from data_processor import process_excel_file


SAMPLE_DATA = [
    {
        "棟別": "K18",
        "基地面積(M2)": 10000,
        "容積率": 0.6,
        "建蔽率": 0.4,
        "開挖深度(M)": 12,
        "耐震係數(gal)": 400,
        "汽車停車位": 30,
        "機車停車位": 60,
        "樓層": [
            {
                "樓層": "1F",
                "狀態": "已成廠",
                "預計成廠年份": "",
                "進駐製程": "測試製程",
                "樓層高度(cm)": 600,
                "無塵室淨高(cm)": 350,
                "樓地板面積(M2)": 5000,
                "無塵室面積(M2)": 1200,
                "生產週邊(M2)": 800,
                "公設(含其他)(公式)(M2)": 1500,
                "廠務設施面積(M2)": {"value": 1500, "details": {"純水": 500, "其他": 1000}},
                "樓層載重kgf/m2": 1000,
            },
            {
                "樓層": "2F",
                "狀態": "未成廠",
                "預計成廠年份": "Y1",
                "進駐製程": "擴建",
                "樓層高度(cm)": 550,
                "無塵室淨高(cm)": 320,
                "樓地板面積(M2)": 4000,
                "無塵室面積(M2)": 1000,
                "生產週邊(M2)": 700,
                "公設(含其他)(公式)(M2)": 1200,
                "廠務設施面積(M2)": {"value": 1100, "details": {"電力": 1100}},
                "樓層載重kgf/m2": 900,
            },
        ],
    }
]


class BuildingDataManagerTests(unittest.TestCase):
    def test_normalize_adds_stable_ids(self):
        first = normalize_dataset(SAMPLE_DATA)
        second = normalize_dataset(SAMPLE_DATA)
        self.assertEqual(first[0]["_building_id"], second[0]["_building_id"])
        self.assertEqual(first[0]["樓層"][0]["_floor_id"], second[0]["樓層"][0]["_floor_id"])
        self.assertEqual(dataset_revision(first), dataset_revision(second))

    def test_normalize_legacy_height_values_to_centimeters(self):
        legacy_data = deepcopy(SAMPLE_DATA)
        floor = legacy_data[0]["樓層"][0]
        floor["樓層高度(cm)"] = "6.0 M"
        floor["無塵室淨高(cm)"] = "3.5m"

        normalized = normalize_dataset(legacy_data)

        self.assertEqual(normalized[0]["樓層"][0]["樓層高度(cm)"], 600)
        self.assertEqual(normalized[0]["樓層"][0]["無塵室淨高(cm)"], 350)

    def test_change_summary_detects_floor_update(self):
        before = normalize_dataset(SAMPLE_DATA)
        after = normalize_dataset(SAMPLE_DATA)
        after[0]["樓層"][0]["樓地板面積(M2)"] = 5200
        summary = summarize_changes(before, after)
        self.assertEqual(summary["floors_updated"], 1)
        self.assertEqual(summary["buildings_updated"], 1)

    def test_readable_workbook_matches_upload_layout_and_round_trips(self):
        workbook_stream = build_readable_workbook(SAMPLE_DATA, [], "tester")
        workbook = load_workbook(workbook_stream)
        self.assertEqual(workbook.sheetnames, ["建物面積總表", "年度新增明細", "異動紀錄", "資料說明"])
        sheet = workbook["建物面積總表"]
        self.assertEqual(sheet["A2"].value, "棟別")
        self.assertEqual(sheet["B2"].value, "樓層")
        workbook.close()

        with tempfile.TemporaryDirectory() as folder:
            input_path = os.path.join(folder, "readable.xlsx")
            cleaned_path = os.path.join(folder, "cleaned.xlsx")
            json_path = os.path.join(folder, "data.json")
            with open(input_path, "wb") as handle:
                handle.write(workbook_stream.getvalue())
            result = process_excel_file(input_path, cleaned_path, json_path)
            self.assertEqual(result["buildings"], 1)
            self.assertEqual(result["rows"], 2)
            with open(json_path, "r", encoding="utf-8") as handle:
                round_trip = json.load(handle)
            self.assertEqual(round_trip[0]["棟別"], "K18")
            self.assertEqual(round_trip[0]["樓層"][1]["樓地板面積(M2)"], 4000)

    def test_standard_workbook_has_fixed_schema_sheets(self):
        workbook_stream = build_standard_workbook(SAMPLE_DATA, "tester")
        workbook = load_workbook(workbook_stream, read_only=True)
        self.assertEqual(workbook.sheetnames, ["metadata", "building_master", "floor_area_detail", "change_log", "data_dictionary"])
        self.assertEqual(workbook["building_master"]["A1"].value, "building_id")
        self.assertEqual(workbook["floor_area_detail"]["J1"].value, "floor_area_m2")
        workbook.close()

        with tempfile.TemporaryDirectory() as folder:
            input_path = os.path.join(folder, "standard.xlsx")
            cleaned_path = os.path.join(folder, "cleaned.xlsx")
            json_path = os.path.join(folder, "data.json")
            with open(input_path, "wb") as handle:
                handle.write(workbook_stream.getvalue())
            result = process_excel_file(input_path, cleaned_path, json_path)
            self.assertEqual(result["input_format"], "standard")
            with open(json_path, "r", encoding="utf-8") as handle:
                round_trip = json.load(handle)
            self.assertEqual(round_trip[0]["棟別"], "K18")
            self.assertEqual(round_trip[0]["樓層"][0]["樓地板面積(M2)"], 5000)


if __name__ == "__main__":
    unittest.main()
